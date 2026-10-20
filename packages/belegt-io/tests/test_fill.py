# Spec: 01 FR-501（AC-501-1/2/3）；04 §2.11 filler-file
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from belegt_io import IOErrorCode, Locator, fill_back, parse_questionnaire


def _answer_for(q):
    nid = q.native_id
    if q.options:
        return True if q.answer_format == "bool" else q.options[-1]
    if q.answer_format == "date":
        return date(2027, 3, 31)
    if q.answer_format == "number":
        if nid == "2.18":
            return Decimal("42.5")  # 字典存 0–100，单元格为百分比格式
        if nid == "5.10":
            return 61
        return Decimal("1234.567") if q.unit_requested else 12
    if q.answer_format == "file":
        return "sha256:" + "b" * 64
    return f"Antwort {nid}"


@pytest.fixture(scope="module")
def filled(oem_path, oem_parsed, tmp_path_factory):
    out = tmp_path_factory.mktemp("fill") / "oem_filled.xlsx"
    answers = {q.locator: _answer_for(q) for q in oem_parsed.questions}
    report = fill_back(oem_path, answers, out, parsed=oem_parsed)
    return out, answers, report


def test_all_written(filled):
    out, answers, report = filled
    assert report.ok, report.to_markdown()
    assert len(report.written) == 100
    assert report.out_sha256 and report.export_input_hash.startswith("sha256:")
    md = report.to_markdown()
    assert "| written |" in md and "2.01" in md


def test_reparse_values_consistent(filled, oem_parsed):
    out, answers, _ = filled
    p2 = parse_questionnaire(out)
    assert len(p2.questions) == 100
    by_id = p2.by_native_id()
    for q in oem_parsed.questions:
        got = by_id[q.native_id].existing_value
        exp = answers[q.locator]
        assert by_id[q.native_id].locator.cell == q.locator.cell
        if q.answer_format == "bool":
            assert got == "Ja"
        elif isinstance(exp, date):
            assert got == datetime(exp.year, exp.month, exp.day)
        elif q.native_id == "2.18":
            assert Decimal(str(got)) == Decimal("0.425")
        elif isinstance(exp, (Decimal, int)):
            assert Decimal(str(got)) == Decimal(exp)
        else:
            assert got == exp


def test_formatting_preserved_and_only_answer_cells_changed(filled, oem_path, oem_parsed):
    out, _, _ = filled
    src, dst = load_workbook(oem_path), load_workbook(out)
    answer_cells = {(q.locator.sheet, q.locator.cell) for q in oem_parsed.questions}
    assert src.sheetnames == dst.sheetnames
    assert dst["Listen"].sheet_state == "hidden"
    assert "JaNein" in dst.defined_names
    for ws in src.worksheets:
        wd = dst[ws.title]
        assert sorted(map(str, ws.merged_cells.ranges)) == sorted(map(str, wd.merged_cells.ranges))
        dvs = sorted((d.type, d.formula1, str(d.sqref)) for d in ws.data_validations.dataValidation)
        dvd = sorted((d.type, d.formula1, str(d.sqref)) for d in wd.data_validations.dataValidation)
        assert dvs == dvd
        assert len(ws.conditional_formatting) == len(wd.conditional_formatting)
        assert ws.freeze_panes == wd.freeze_panes
        for row in ws.iter_rows():
            for c in row:
                d = wd[c.coordinate]
                assert c.number_format == d.number_format or (ws.title, c.coordinate) in answer_cells
                assert c.fill.fgColor.rgb == d.fill.fgColor.rgb
                assert bool(c.font.b) == bool(d.font.b)
                if (ws.title, c.coordinate) not in answer_cells:
                    assert c.value == d.value, (ws.title, c.coordinate)
    ws = dst["2 Umwelt & Energie"]
    assert str(ws["G2"].value).startswith("=COUNTA(")  # 公式保留
    assert dst["Übersicht"]["B4"].value == "='2 Umwelt & Energie'!G2"
    assert ws["C25"].number_format == "0.0%"


def test_deterministic_output(filled, oem_path, oem_parsed, tmp_path):
    out, answers, report = filled
    again = fill_back(oem_path, answers, tmp_path / "again.xlsx", parsed=oem_parsed)
    assert again.out_sha256 == report.out_sha256
    assert again.export_input_hash == report.export_input_hash
    # 键形式不影响 export_input_hash
    by_str = {f"{k.sheet}!{k.cell}": v for k, v in answers.items()}
    assert fill_back(oem_path, by_str, tmp_path / "s.xlsx", parsed=oem_parsed).export_input_hash == report.export_input_hash


def test_rejections(oem_path, oem_parsed, tmp_path):
    q = oem_parsed.by_native_id()
    answers = {
        q["2.06"].locator: "Gibt es nicht",  # 非法枚举
        "1 Unternehmen!B6": "Überschreiben",  # 题干单元格（非答案）
        "2 Umwelt & Energie!G2": 1,  # 公式单元格
        q["5.10"].locator: 150,  # whole 0–100 校验
        q["2.01"].locator: "viel",  # 数字题给文本
        Locator(blob_sha256="sha256:" + "0" * 64, sheet="1 Unternehmen", cell="C6"): "x",  # 他人文件
        q["1.01"].locator: "Musterfirma GmbH",
    }
    rep = fill_back(oem_path, answers, tmp_path / "r.xlsx", parsed=oem_parsed)
    codes = sorted(c.code for c in rep.rejected)
    assert codes == sorted([IOErrorCode.ENUM_INVALID, IOErrorCode.NOT_ANSWER_CELL, IOErrorCode.NOT_ANSWER_CELL,
                            IOErrorCode.TYPE_MISMATCH, IOErrorCode.TYPE_MISMATCH, IOErrorCode.HASH_MISMATCH])
    assert len(rep.written) == 1
    wb = load_workbook(tmp_path / "r.xlsx")
    assert wb["1 Unternehmen"]["B6"].value == "Vollständiger Firmenname (laut Handelsregister)"
    assert wb["2 Umwelt & Energie"]["C11"].value is None
    assert wb["1 Unternehmen"]["C6"].value == "Musterfirma GmbH"


def test_number_rounding_and_german_string(oem_path, oem_parsed, tmp_path):
    q = oem_parsed.by_native_id()
    rep = fill_back(oem_path, {q["2.12"].locator: {"value": Decimal("4820.04449"), "round": 2},
                               q["2.13"].locator: "1.234,5"}, tmp_path / "n.xlsx", parsed=oem_parsed)
    assert rep.ok
    ws = load_workbook(tmp_path / "n.xlsx")["2 Umwelt & Energie"]
    assert Decimal(str(ws[q["2.12"].locator.cell].value)) == Decimal("4820.04")
    assert Decimal(str(ws[q["2.13"].locator.cell].value)) == Decimal("1234.5")


def test_merged_answer_cell_redirects_to_top_left(oem_path, oem_parsed, tmp_path):
    rep = fill_back(oem_path, {"1 Unternehmen!D7": "Musterstraße 1, 12345 Musterstadt"}, tmp_path / "m.xlsx",
                    parsed=oem_parsed)
    assert rep.ok and rep.written[0].cell == "C7"
    assert load_workbook(tmp_path / "m.xlsx")["1 Unternehmen"]["C7"].value.startswith("Musterstraße")


def test_csv_fill_roundtrip(ex_dir, tmp_path):
    src = ex_dir / "kurz.csv"
    p = parse_questionnaire(src)
    q = p.by_native_id()
    rep = fill_back(src, {q["3"].locator: Decimal("1835.2"), q["7"].locator: True, q["8"].locator: date(2027, 5, 31),
                          q["1"].locator: "Musterfirma GmbH"}, tmp_path / "k.csv", parsed=p)
    assert rep.ok and len(rep.written) == 4
    text = (tmp_path / "k.csv").read_text(encoding="utf-8-sig")
    assert "1835,2" in text and "31.05.2027" in text and ";Ja;" in text
    p2 = parse_questionnaire(tmp_path / "k.csv")
    assert p2.by_native_id()["3"].existing_value == "1835,2"
    assert len(p2.questions) == 15
