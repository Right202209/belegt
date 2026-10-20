# Spec: 05 §4.6（PACT v3.0.3）、§6.9（CBAM communication）；02 §4.3.3/§4.3.7/§4.3.8；01 AC-503-1、AC-503-2
from __future__ import annotations

import json
from datetime import date
from decimal import Decimal

import pytest
from jsonschema import Draft202012Validator
from openpyxl import load_workbook

from belegt_io import parse_questionnaire
from belegt_io.examples import sample_cbam, sample_cmrt, sample_pcf_input
from belegt_io.export import (
    BatteryCfInput, ExportValidationError, LifecycleStages, export_battery_passport_cf, export_catenax_pcf,
    export_cbam_communication, export_cmrt, export_pact, pcf_input_from_datapoints, validate_pact,
)
from belegt_io.export.common import load_schema, loads_json
from belegt_io.models import Period


# —— PACT ——

def test_pact_export_passes_schema_and_rules():
    res = export_pact(sample_pcf_input())
    doc = res.document
    assert doc["specVersion"] == "3.0.3"
    pcf = doc["pcf"]
    # 05 §4.6：End 为开区间 = 闭区间 end + 1 天
    assert pcf["referencePeriodStart"] == "2025-01-01T00:00:00Z"
    assert pcf["referencePeriodEnd"] == "2026-01-01T00:00:00Z"
    assert pcf["declaredUnitOfMeasurement"] == "kilogram"
    assert pcf["pcfExcludingBiogenicUptake"] == "2.4137" and isinstance(pcf["primaryDataShare"], str)
    assert pcf["ipccCharacterizationFactors"] == ["AR6"] and pcf["crossSectoralStandards"] == ["ISO14067"]
    # 独立再验一次（官方 OpenAPI 3.0.3 派生 schema）
    Draft202012Validator(load_schema("pact-productfootprint-3.0.3.schema.json")).validate(json.loads(res.to_json()))
    errors, _ = validate_pact(json.loads(res.to_json()))
    assert errors == []
    required_cf = load_schema("pact-productfootprint-3.0.3.schema.json")["$defs"]["CarbonFootprint"]["required"]
    assert set(required_cf) <= set(pcf)
    assert res.export_input_hash == export_pact(sample_pcf_input()).export_input_hash


@pytest.mark.parametrize("change,needle", [
    ({"exempted_emissions_percent": Decimal("4")}, "PACT-R-EXEMPT"),
    ({"primary_data_share": None}, "primaryDataShare"),
    ({"company_ids": ["not-a-urn"]}, "companyIds"),
    ({"declared_unit": "barrel"}, "declaredUnitOfMeasurement"),
    ({"product_mass_per_declared_unit": Decimal("0")}, "PACT-R-MASS"),
    ({"geography_region_or_subregion": "Europe", "geography_country": None}, None),
])
def test_pact_rule_violations(change, needle):
    data = sample_pcf_input().model_dump()
    data.update(change)
    if needle is None:
        assert export_pact(data).document["pcf"]["geographyRegionOrSubregion"] == "Europe"
        return
    with pytest.raises(ExportValidationError) as e:
        export_pact(data)
    assert any(needle in x for x in e.value.errors), e.value.errors


def test_pact_dqi_mandatory_from_2027():
    data = sample_pcf_input().model_dump()
    data["dqi"] = None
    assert any("dqi" in w for w in export_pact(data).warnings)
    data["created"] = "2027-01-05T00:00:00Z"
    with pytest.raises(ExportValidationError):
        export_pact(data)


def test_pcf_from_datapoints():
    base = sample_pcf_input().model_dump(exclude={"pcf_excluding_biogenic_uptake", "pcf_including_biogenic_uptake",
                                                   "declared_unit", "product_mass_per_declared_unit",
                                                   "reference_period", "primary_data_share"})
    inp = pcf_input_from_datapoints({
        "dp:prod.pcf.cradle_to_gate.excluding_biogenic@1.0.0": Decimal("2.4137"),
        "dp:prod.pcf.cradle_to_gate@1.0.0": Decimal("2.4137"),
        "dp:prod.pcf.declared_unit": "kg",
        "dp:prod.mass.net_per_unit": Decimal("1"),
        "dp:prod.pcf.reference_period": {"start": "2025-01-01", "end": "2025-12-31"},
        "dp:prod.pcf.primary_data_share": Decimal("38.5"),
    }, **base)
    assert export_pact(inp).document == export_pact(sample_pcf_input()).document


# —— Catena-X ——

def test_catenax_export_validates_and_uses_numbers():
    res = export_catenax_pcf(sample_pcf_input(), {"product_or_sector_specific_rules": ["Catena-X PCF Rulebook"]})
    text = res.to_json()
    assert '"pcfExcludingBiogenicUptake": 2.4137' in text  # JSON number，精确输出
    doc = loads_json(text)
    prod = doc["productLifeCycleStagesAndEmissions"][0]["productionStage"][0]
    assert prod["pcfIncludingBiogenicUptake"] == Decimal("2.4137")
    geo = doc["pcfAssessmentAndMethodology"][0]["pcfAssessmentInformation"][0]["geography"][0]
    assert geo == {"geographyRegionOrSubregion": "Western Europe", "geographyCountry": "DE"}
    assert doc["scopeOfPcfForm"][0]["specVersion"] == "urn:samm:io.catenax.pcf:10.0.0"
    std = doc["pcfAssessmentAndMethodology"][0]["pcfMethodology"][0]["standards"][0]
    assert std["crossSectoralStandards"] == "ISO 14067"


def test_catenax_requires_validity_period():
    data = sample_pcf_input().model_dump()
    data["validity_period"] = None
    with pytest.raises(ExportValidationError) as e:
        export_catenax_pcf(data)
    assert any("validityPeriodEnd" in x for x in e.value.errors)


# —— CBAM communication ——

def test_cbam_export_roundtrip(tmp_path):
    out = tmp_path / "cbam.xlsx"
    res = export_cbam_communication(sample_cbam(), out)
    assert res.out_path == str(out) and res.export_input_hash.startswith("sha256:")
    p = parse_questionnaire(out)
    assert p.template_ref == "tm:cbam_communication@2024"
    vals = {}
    for q in p.questions:
        if q.existing_value is not None:
            for h in q.exact_mappings:
                vals.setdefault(h.datapoint_id, []).append(q.existing_value)
    assert vals["dp:prod.cbam.installation_id"] == ["CN-TJ-STEEL-0815"]
    assert vals["dp:prod.identity.cn_code"] == ["72091610", "72104900"]
    assert [Decimal(str(v)) for v in vals["dp:prod.cbam.see_direct"]] == [Decimal("2.041234"), Decimal("2.1123")]
    assert vals["dp:prod.cbam.production_route"] == ["Blast furnace route"]
    ws = load_workbook(out)["Summary_Products"]
    assert ws["E4"].number_format == "0.000000"
    # 确定性
    res2 = export_cbam_communication(sample_cbam(), tmp_path / "cbam2.xlsx")
    assert res2.export_input_hash == res.export_input_hash
    assert out.read_bytes() == (tmp_path / "cbam2.xlsx").read_bytes()


def test_cbam_validation_errors(tmp_path):
    data = sample_cbam().model_dump()
    data["goods"][0]["cn_code"] = "7209"
    with pytest.raises(Exception):
        export_cbam_communication(data, tmp_path / "x.xlsx")
    data = sample_cbam().model_dump()
    data["goods"][0]["see_direct"] = Decimal("-1")
    with pytest.raises(Exception):
        export_cbam_communication(data, tmp_path / "x.xlsx")


def test_cbam_into_given_template(tmp_path, ex_dir):
    out = tmp_path / "filled.xlsx"
    export_cbam_communication(sample_cbam(), out, template_path=ex_dir / "cbam.xlsx")
    assert load_workbook(out)["A_InstData"]["C3"].value == "CN-TJ-STEEL-0815"


# —— CMRT ——

def test_cmrt_export_roundtrip(tmp_path):
    out = tmp_path / "cmrt.xlsx"
    export_cmrt(sample_cmrt(), out)
    p = parse_questionnaire(out)
    q = p.by_native_id()
    assert q["Q1.Tin"].existing_value == "Yes" and q["Q1.Gold"].existing_value == "No"
    assert q["B"].existing_value == "https://musterfirma.example/policy"
    ws = load_workbook(out)["Smelter List"]
    assert [ws[f"{c}4"].value for c in "ABCD"] == ["Tin", "Demo Tin Smelter", "Indonesia", "CID000001"]


def test_cmrt_validation(tmp_path):
    data = sample_cmrt().model_dump()
    data["smelters"][0]["smelter_id"] = "X123"
    with pytest.raises(Exception):
        export_cmrt(data, tmp_path / "x.xlsx")
    data = sample_cmrt().model_dump()
    data["company_answers"]["A"] = "Vielleicht"
    with pytest.raises(Exception):
        export_cmrt(data, tmp_path / "x.xlsx")


# —— 电池护照 ——

def _battery(**kw):
    base = dict(manufacturer_name="Demo Cells GmbH", battery_model_id="DC-NMC-77", manufacturing_plant_country="DE",
                carbon_footprint_per_kwh=Decimal("61.4"), chemistry="NMC", created="2026-10-19T00:00:00Z",
                lifecycle_stages=LifecycleStages(raw_material_acquisition_and_pre_processing=Decimal("44.1"),
                                                 main_product_production=Decimal("14.2"), distribution=Decimal("1.3"),
                                                 end_of_life_and_recycling=Decimal("1.8")))
    base.update(kw)
    return BatteryCfInput(**base)


def test_battery_upstream_fragment():
    res = export_battery_passport_cf(_battery())
    cf = res.document["batteryCarbonFootprint"]
    assert cf["unit"] == "kgCO2e/kWh" and cf["carbonFootprintPerFunctionalUnit"] == Decimal("61.4")
    assert cf["declarationStatus"] == "not_a_declaration"
    assert '"carbonFootprintPerFunctionalUnit": 61.4' in res.to_json()


def test_battery_passport_mode_excludes_declaration():
    res = export_battery_passport_cf(_battery(), mode="passport")
    assert "batteryCarbonFootprint" not in res.document  # AC-503-2
    assert "batteryCarbonFootprint" in res.omitted_fields and "not yet adopted" in res.omitted_fields["batteryCarbonFootprint"]
    assert res.document["batteryChemistry"] == "nmc"


def test_battery_stage_sum_checked():
    bad = _battery(carbon_footprint_per_kwh=Decimal("70"))
    with pytest.raises(ExportValidationError):
        export_battery_passport_cf(bad)
