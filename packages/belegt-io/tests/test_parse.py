# Spec: 03 §3.3 Question；04 §2.3；01 AC-501-1（合并单元格、下拉）
from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from belegt_io import BelegtIOError, IOErrorCode, Locator, file_sha256, parse_questionnaire
from belegt_io.workbooks import OEM_SHEETS


def test_oem_has_exactly_100_questions(oem_parsed):
    assert len(oem_parsed.questions) == 100
    assert [q.ordinal for q in oem_parsed.questions] == list(range(1, 101))
    ids = [q.native_id for q in oem_parsed.questions]
    assert len(set(ids)) == 100
    expected = [nid for _, secs in OEM_SHEETS for _, qs in secs for nid, *_ in qs]
    assert ids == expected


def test_locators_point_to_answer_cells(oem_parsed, oem_path):
    sha = file_sha256(oem_path)
    assert oem_parsed.blob_sha256 == sha
    wb = load_workbook(oem_path)
    for q in oem_parsed.questions:
        loc = q.locator
        assert loc.blob_sha256 == sha
        assert loc.mime.endswith("spreadsheetml.sheet")
        ws = wb[loc.sheet]
        row = int(loc.cell[1:])
        assert loc.cell.startswith("C"), loc
        assert ws[f"A{row}"].value == q.native_id
        assert ws[f"B{row}"].value == next(iter(q.text.values()))
        assert q.text_locator.cell == f"B{row}"
        assert str(loc) == f"{sha}#{loc.sheet}!{loc.cell}"
    q201 = oem_parsed.by_native_id()["2.01"]
    assert q201.locator.sheet == "2 Umwelt & Energie" and q201.locator.cell == "C6"


def test_formats_units_periods(oem_parsed):
    q = oem_parsed.by_native_id()
    assert q["2.01"].answer_format == "number" and q["2.01"].unit_requested == "tCO2e"
    assert q["2.01"].period_requested.start == date(2025, 1, 1) and q["2.01"].period_requested.end == date(2025, 12, 31)
    assert q["2.22"].unit_requested == "m3"
    assert q["2.29"].unit_requested == "kgCO2e/kg"
    assert q["3.06"].unit_requested == "tCO2e/t"
    assert q["4.03"].unit_requested == "%w/w"
    assert q["1.10"].period_type_hint == "point_in_time" and q["1.10"].period_requested.start == date(2025, 12, 31)
    assert q["1.11"].period_type_hint == "fiscal_year" and q["1.11"].period_requested is None
    assert q["5.02"].answer_format == "date"
    assert q["5.03"].answer_format == "file"
    assert q["5.10"].answer_format == "number"  # whole-number data validation
    assert all(x.language == "de" for x in oem_parsed.questions)
    assert oem_parsed.language == "de"


def test_dropdowns_become_enum_candidates(oem_parsed):
    q = oem_parsed.by_native_id()
    assert q["1.06"].answer_format == "bool" and q["1.06"].options == ["Ja", "Nein"]  # 定义名称 JaNein
    assert q["1.06"].enum_source == "data_validation:JaNein"
    assert q["1.05"].options == ["GmbH", "AG", "KG", "GmbH & Co. KG", "SE", "Sonstige"]  # 跨表区域
    assert q["2.06"].answer_format == "enum" and q["2.06"].options[0] == "Operative Kontrolle"  # 字面量
    assert q["4.06"].answer_format == "enum" and "Unbekannt" in q["4.06"].options


def test_merged_cells_and_sections(oem_parsed):
    q = oem_parsed.by_native_id()
    assert q["1.02"].hints.get("answer_cell_merged") is True
    assert q["1.02"].locator.cell == "C7"
    assert q["2.01"].section == "Treibhausgasemissionen"
    assert q["2.12"].section == "Energie"
    skipped = {s.sheet: s.skipped_reason for s in oem_parsed.sheets}
    assert skipped["Listen"] == "hidden_sheet"
    assert skipped["Hinweise"] == "no_header_detected"


def test_not_a_known_template(oem_parsed):
    assert oem_parsed.template_ref is None
    assert all(not q.exact_mappings for q in oem_parsed.questions)


def test_question_fields_match_contract(oem_parsed):
    d = oem_parsed.questions[0].to_question_fields("req_01JD8Y3K5M2N4P6Q8R0S2T4V6W")
    assert set(d) <= {"request_id", "ordinal", "native_id", "text", "answer_format", "options", "unit_requested",
                      "period_requested", "locator", "parent_question_id", "status"}
    assert d["status"] == "parsed" and d["locator"]["cell"] == "C6" and d["locator"]["sheet"] == "1 Unternehmen"


def test_csv(ex_dir):
    p = parse_questionnaire(ex_dir / "kurz.csv")
    assert len(p.questions) == 15
    q = p.by_native_id()
    assert q["3"].locator.sheet is None and q["3"].locator.cell == "C4"
    assert q["3"].unit_requested == "tCO2e" and q["3"].period_requested.start == date(2025, 1, 1)
    assert q["7"].answer_format == "bool"
    assert p.language == "de"


def test_locator_parse_roundtrip():
    sha = "sha256:" + "a" * 64
    for s in (f"{sha}#Sheet 1!B12", "'My Sheet'!$C$5", "D7", f"{sha}#p3"):
        loc = Locator.parse(s)
        assert Locator.parse(str(loc)) == loc or not loc.blob_sha256
    assert Locator.parse("'My Sheet'!$C$5").sheet == "My Sheet"
    assert Locator.parse(f"{sha}#p3").page == 3


def test_errors(tmp_path):
    bad = tmp_path / "x.pdf"
    bad.write_bytes(b"%PDF")
    with pytest.raises(BelegtIOError) as e:
        parse_questionnaire(bad)
    assert e.value.code == IOErrorCode.UNSUPPORTED_FORMAT
    enc = tmp_path / "locked.xlsx"
    enc.write_bytes(b"\xd0\xcf\x11\xe0 not a zip (encrypted OLE container)")
    with pytest.raises(BelegtIOError) as e:
        parse_questionnaire(enc)
    assert e.value.code == IOErrorCode.FILE_UNREADABLE
