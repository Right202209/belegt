# Spec: 01 §2.1（OEM 问卷旅程）、AC-501-1；02 §4.3.4（CBAM 沟通模板）、§4.3.7（CMRT）
"""用 openpyxl 程序化生成示例/模板工作簿：OEM 德语 100 题问卷、CBAM 沟通模板、简化 CMRT、CSV、示例 PDF。

CBAM 与 CMRT 生成的是 **Belegt 样例版面**（工作表名/列名按 02 §4.3 映射文件，官方单元格位置 [需核对]），
不是欧盟委员会 / RMI 的官方文件；题干为本包自拟短句，不复制任何受版权保护的问卷原文。
"""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

FIXED_CREATED = "2026-10-01T00:00:00Z"
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _fix_props(wb: Workbook) -> None:
    from datetime import datetime, timezone

    ts = datetime(2026, 10, 1, tzinfo=timezone.utc).replace(tzinfo=None)
    wb.properties.created = ts
    wb.properties.modified = ts
    wb.properties.creator = "belegt-io examples"


def _save(wb: Workbook, path: str | Path) -> Path:
    from .fill import _deterministic_rezip

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    _fix_props(wb)
    wb.save(p)
    _deterministic_rezip(p, FIXED_CREATED)
    return p


# —— OEM 德语 100 题 ——
# (Nr, Frage, Format, Einheit, Optionen-Schlüssel)  Format: n=Zahl, t=Text, b=Ja/Nein, d=Datum, e=Enum, f=Datei, p=Prozent
OEM_SHEETS: list[tuple[str, list[tuple[str, list[tuple[str, str, str, str, str | None]]]]]] = [
    ("1 Unternehmen", [
        ("Allgemeine Angaben", [
            ("1.01", "Vollständiger Firmenname (laut Handelsregister)", "t", "", None),
            ("1.02", "Anschrift des Hauptsitzes", "tm", "", None),
            ("1.03", "Umsatzsteuer-Identifikationsnummer", "t", "", None),
            ("1.04", "DUNS-Nummer", "t", "", None),
            ("1.05", "Rechtsform", "e", "", "rechtsform"),
            ("1.06", "Gehört Ihr Unternehmen zu einem Konzern?", "b", "", None),
            ("1.07", "Name der Muttergesellschaft (falls zutreffend)", "t", "", None),
            ("1.08", "Haupttätigkeit nach NACE-Code", "t", "", None),
            ("1.09", "Webseite des Unternehmens", "t", "", None),
        ]),
        ("Kennzahlen", [
            ("1.10", "Anzahl Mitarbeitende zum 31.12.2025", "n", "Anzahl", None),
            ("1.11", "Umsatz im Geschäftsjahr 2025", "n", "EUR", None),
            ("1.12", "Anzahl der Produktionsstandorte", "n", "Anzahl", None),
        ]),
        ("Ansprechpartner", [
            ("1.13", "Name der verantwortlichen Person für Nachhaltigkeit", "t", "", None),
            ("1.14", "E-Mail-Adresse der verantwortlichen Person", "t", "", None),
            ("1.15", "Telefonnummer der verantwortlichen Person", "t", "", None),
        ]),
    ]),
    ("2 Umwelt & Energie", [
        ("Treibhausgasemissionen", [
            ("2.01", "Wie hoch waren Ihre Scope-1-Emissionen im Berichtsjahr?", "n", "tCO2e", None),
            ("2.02", "Wie hoch waren Ihre Scope-2-Emissionen (standortbasiert)?", "n", "tCO2e", None),
            ("2.03", "Wie hoch waren Ihre Scope-2-Emissionen (marktbasiert)?", "n", "tCO2e", None),
            ("2.04", "Erfassen Sie Scope-3-Emissionen?", "b", "", None),
            ("2.05", "Scope-3-Emissionen Kategorie 1 (eingekaufte Waren und Dienstleistungen)", "n", "tCO2e", None),
            ("2.06", "Nach welchem Konsolidierungsansatz bilanzieren Sie?", "e", "", "konsolidierung"),
            ("2.07", "Basisjahr Ihrer Treibhausgasbilanz", "n", "Jahr", None),
        ]),
        ("Klimaziele", [
            ("2.08", "Hat Ihr Unternehmen ein Reduktionsziel für Treibhausgase?", "b", "", None),
            ("2.09", "Zieljahr des Reduktionsziels", "n", "Jahr", None),
            ("2.10", "Angestrebte Reduktion gegenüber dem Basisjahr", "n", "%", None),
            ("2.11", "Ist Ihr Ziel durch die Science Based Targets initiative validiert?", "e", "", "sbti"),
        ]),
        ("Energie", [
            ("2.12", "Gesamtenergieverbrauch im Berichtsjahr", "n", "MWh", None),
            ("2.13", "Stromverbrauch im Berichtsjahr", "n", "MWh", None),
            ("2.14", "Erdgasverbrauch im Berichtsjahr", "n", "MWh", None),
            ("2.15", "Heizölverbrauch im Berichtsjahr", "n", "L", None),
            ("2.16", "Dieselverbrauch der eigenen Fahrzeugflotte", "n", "L", None),
            ("2.17", "Bezug von Fernwärme", "n", "MWh", None),
            ("2.18", "Anteil erneuerbarer Energien am Stromverbrauch", "p", "%", None),
            ("2.19", "Eigenerzeugung aus Photovoltaik", "n", "kWh", None),
            ("2.20", "Beziehen Sie Ökostrom über einen eigenen Liefervertrag?", "b", "", None),
            ("2.21", "Menge der entwerteten Herkunftsnachweise", "n", "MWh", None),
        ]),
        ("Wasser und Abfall", [
            ("2.22", "Wasserentnahme gesamt", "n", "m³", None),
            ("2.23", "Abfallaufkommen gesamt", "n", "t", None),
            ("2.24", "Davon gefährliche Abfälle", "n", "t", None),
            ("2.25", "Recyclingquote der Abfälle", "n", "%", None),
        ]),
        ("Umweltmanagement und Produkt-CO2", [
            ("2.26", "Beschreiben Sie kurz Ihr Umweltmanagementsystem.", "tm", "", None),
            ("2.27", "Datum des letzten Energieaudits", "d", "", None),
            ("2.28", "Liegt für das gelieferte Produkt ein Product Carbon Footprint (PCF) vor?", "b", "", None),
            ("2.29", "PCF des gelieferten Produkts (cradle-to-gate)", "n", "kgCO2e/kg", None),
            ("2.30", "Bezugsjahr des PCF", "n", "Jahr", None),
        ]),
    ]),
    ("3 CBAM", [
        ("CBAM-Waren", [
            ("3.01", "Liefern Sie Waren, die unter den CBAM fallen (Eisen/Stahl, Aluminium)?", "b", "", None),
            ("3.02", "KN-Codes (8-stellig) der gelieferten CBAM-Waren", "t", "", None),
            ("3.03", "Kennung der Produktionsanlage (Installation ID)", "t", "", None),
            ("3.04", "Ursprungsland der Produktion (ISO-Code)", "t", "", None),
            ("3.05", "Produktionsroute", "e", "", "route"),
        ]),
        ("Graue Emissionen", [
            ("3.06", "Spezifische direkte graue Emissionen (SEE direkt)", "n", "tCO2e/t", None),
            ("3.07", "Spezifische indirekte graue Emissionen (SEE indirekt)", "n", "tCO2e/t", None),
            ("3.08", "Spezifischer Stromverbrauch", "n", "MWh/t", None),
            ("3.09", "Wurden Standardwerte statt tatsächlicher Werte verwendet?", "e", "", "standardwerte"),
            ("3.10", "Liegt ein Prüfbericht einer akkreditierten Prüfstelle vor?", "b", "", None),
            ("3.11", "Name der Prüfstelle", "t", "", None),
            ("3.12", "Datum der Verifizierung", "d", "", None),
            ("3.13", "Im Ursprungsland tatsächlich gezahlter CO2-Preis", "n", "EUR", None),
            ("3.14", "Gelieferte Menge CBAM-Waren im Berichtsjahr", "n", "t", None),
            ("3.15", "E-Mail des Anlagenbetreibers für CBAM-Rückfragen", "t", "", None),
        ]),
    ]),
    ("4 REACH & Materialien", [
        ("Stoffe", [
            ("4.01", "Enthält das Produkt SVHC-Stoffe der Kandidatenliste über 0,1 Massen-%?", "b", "", None),
            ("4.02", "Bezeichnung und CAS-Nummer der enthaltenen SVHC-Stoffe", "tm", "", None),
            ("4.03", "Konzentration des SVHC-Stoffes im Erzeugnis", "n", "%w/w", None),
            ("4.04", "SCIP-Nummer der Meldung", "t", "", None),
            ("4.05", "Ist das Produkt RoHS-konform?", "b", "", None),
            ("4.06", "Enthält das Produkt PFAS?", "e", "", "jnu"),
        ]),
        ("Mineralien und Materialdaten", [
            ("4.07", "Liegt ein aktuelles CMRT für Konfliktmineralien vor?", "e", "", "cmrt"),
            ("4.08", "Liegt ein EMRT (Kobalt, Glimmer) vor?", "b", "", None),
            ("4.09", "Sind die Materialdaten im IMDS hinterlegt?", "b", "", None),
            ("4.10", "IMDS-ID des Datenblatts", "t", "", None),
            ("4.11", "Rezyklatanteil im Produkt", "n", "%", None),
            ("4.12", "Nettomasse des Produkts je Stück", "n", "kg", None),
        ]),
        ("Verpackung", [
            ("4.13", "Hauptverpackungsmaterial", "e", "", "verpackung"),
            ("4.14", "Verpackungsgewicht je Liefereinheit", "n", "kg", None),
            ("4.15", "Ist die Verpackung ein Mehrwegsystem?", "b", "", None),
        ]),
    ]),
    ("5 Zertifikate", [
        ("Managementsysteme", [
            ("5.01", "Ist Ihr Standort nach ISO 14001 zertifiziert?", "b", "", None),
            ("5.02", "ISO 14001: gültig bis", "d", "", None),
            ("5.03", "Bitte laden Sie das ISO-14001-Zertifikat hoch.", "f", "", None),
            ("5.04", "Ist Ihr Standort nach ISO 50001 zertifiziert?", "b", "", None),
            ("5.05", "ISO 50001: gültig bis", "d", "", None),
            ("5.06", "Ist Ihr Standort nach IATF 16949 zertifiziert?", "b", "", None),
            ("5.07", "IATF 16949: gültig bis", "d", "", None),
            ("5.08", "Ist Ihr Standort nach ISO 45001 zertifiziert?", "b", "", None),
        ]),
        ("Ratings", [
            ("5.09", "Aktuelle EcoVadis-Medaille", "e", "", "ecovadis"),
            ("5.10", "EcoVadis-Gesamtpunktzahl", "w", "Punkte", None),
        ]),
    ]),
    ("6 Lieferkette & Soziales", [
        ("Lieferkette", [
            ("6.01", "Haben Sie einen Verhaltenskodex für Lieferanten?", "b", "", None),
            ("6.02", "Führen Sie eine Risikoanalyse nach LkSG durch?", "b", "", None),
            ("6.03", "Gibt es einen Beschwerdemechanismus für externe Hinweisgeber?", "b", "", None),
            ("6.04", "Anzahl direkter Lieferanten", "n", "Anzahl", None),
            ("6.05", "Anteil der nach Nachhaltigkeitskriterien bewerteten Lieferanten", "n", "%", None),
        ]),
        ("Arbeitssicherheit und Soziales", [
            ("6.06", "Anzahl meldepflichtiger Arbeitsunfälle im Berichtsjahr", "n", "Anzahl", None),
            ("6.07", "Unfallhäufigkeitsrate (LTIFR)", "n", "", None),
            ("6.08", "Durchschnittliche Weiterbildungsstunden je Mitarbeitenden", "n", "h", None),
            ("6.09", "Frauenanteil in Führungspositionen", "n", "%", None),
            ("6.10", "Zahlen Sie an allen Standorten mindestens den gesetzlichen Mindestlohn?", "b", "", None),
            ("6.11", "Gibt es eine Richtlinie gegen Kinder- und Zwangsarbeit?", "b", "", None),
        ]),
        ("Governance", [
            ("6.12", "Gibt es eine Antikorruptionsrichtlinie?", "b", "", None),
            ("6.13", "Höhe der Bußgelder wegen Korruption im Berichtsjahr", "n", "EUR", None),
            ("6.14", "Verfügen Sie über ein Hinweisgebersystem?", "b", "", None),
            ("6.15", "Link zum aktuellen Nachhaltigkeitsbericht", "t", "", None),
        ]),
    ]),
]

ENUMS_LITERAL = {
    "konsolidierung": "Operative Kontrolle,Finanzielle Kontrolle,Equity-Share",
    "sbti": "Ja – validiert,Commitment,Nein",
    "route": "Hochofen,Elektrolichtbogenofen,DRI-EAF,Sonstige",
    "standardwerte": "Tatsächliche Werte,Standardwerte,Gemischt",
    "jnu": "Ja,Nein,Unbekannt",
    "cmrt": "Ja,Nein,In Bearbeitung",
    "verpackung": "Kunststoff,Karton,Holz,Metall,Mehrweg",
    "ecovadis": "Platin,Gold,Silber,Bronze,Keine",
}


def build_oem_questionnaire_de(path: str | Path) -> Path:
    """OEM-Lieferantenfragebogen（德语，100 题，6 个题目工作表 + 说明/概览/隐藏列表页）。"""
    wb = Workbook()
    hinweise = wb.active
    hinweise.title = "Hinweise"
    hinweise["A1"] = "Lieferanten-Nachhaltigkeitsfragebogen 2026 – Musterkunde Automotive AG"
    hinweise["A1"].font = TITLE_FONT
    hinweise["A3"] = ("Bitte beantworten Sie alle gelb markierten Felder in den Registerkarten 1 bis 6 vollständig "
                      "und senden Sie die Datei bis zum 15.12.2026 zurück.")
    hinweise["A4"] = ("Zahlenangaben bitte ohne Einheit in die gelben Felder eintragen; die erwartete Einheit steht "
                      "jeweils rechts daneben. Für Rückfragen steht Ihnen der Einkauf zur Verfügung.")
    hinweise.column_dimensions["A"].width = 120

    listen = wb.create_sheet("Listen")
    for i, v in enumerate(["Ja", "Nein"], start=1):
        listen.cell(row=i, column=1, value=v)
    for i, v in enumerate(["GmbH", "AG", "KG", "GmbH & Co. KG", "SE", "Sonstige"], start=1):
        listen.cell(row=i, column=3, value=v)
    listen.sheet_state = "hidden"
    wb.defined_names["JaNein"] = DefinedName("JaNein", attr_text="Listen!$A$1:$A$2")

    summary = wb.create_sheet("Übersicht")
    summary["A1"] = "Bearbeitungsstand (automatisch berechnet)"
    summary["A1"].font = TITLE_FONT

    total = 0
    for s_idx, (sheet_name, sections) in enumerate(OEM_SHEETS, start=1):
        ws = wb.create_sheet(sheet_name)
        ws.merge_cells("A1:E1")
        ws["A1"] = f"Lieferanten-Nachhaltigkeitsfragebogen 2026 – {sheet_name}"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = "Berichtsjahr: 2025 · Bitte nur gelbe Felder ausfüllen."
        ws["G1"] = "Fortschritt"
        ws["G1"].font = Font(bold=True)
        headers = ["Nr.", "Frage", "Antwort", "Einheit", "Kommentar / Nachweis"]
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=c, value=h)
            cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
        widths = {"A": 8, "B": 70, "C": 28, "D": 12, "E": 36, "G": 14}
        for k, w in widths.items():
            ws.column_dimensions[k].width = w
        ws.freeze_panes = "C5"
        dv_janein = DataValidation(type="list", formula1="JaNein", allow_blank=True)
        dv_rechts = DataValidation(type="list", formula1="Listen!$C$1:$C$6", allow_blank=True)
        dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)  # > 2000-01-01
        dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
        dv_whole = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=True)
        dv_enums: dict[str, DataValidation] = {}
        for dv in (dv_janein, dv_rechts, dv_date, dv_num, dv_whole):
            dv.error = "Ungültige Eingabe"
            ws.add_data_validation(dv)
        r = 5
        first_q = r
        for section, qs in sections:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            sc = ws.cell(row=r, column=1, value=section)
            sc.font, sc.fill = Font(bold=True), SECTION_FILL
            r += 1
            for nid, text, fmt, unit, opt in qs:
                total += 1
                ws.cell(row=r, column=1, value=nid).border = BORDER
                qc = ws.cell(row=r, column=2, value=text)
                qc.alignment, qc.border = WRAP, BORDER
                ac = ws.cell(row=r, column=3)
                ac.fill, ac.border = INPUT_FILL, BORDER
                if unit:
                    ws.cell(row=r, column=4, value=unit).border = BORDER
                ws.cell(row=r, column=5).border = BORDER
                ref = ac.coordinate
                if fmt == "b":
                    dv_janein.add(ref)
                elif fmt == "e" and opt == "rechtsform":
                    dv_rechts.add(ref)
                elif fmt == "e":
                    if opt not in dv_enums:
                        dv_enums[opt] = DataValidation(type="list", formula1=f'"{ENUMS_LITERAL[opt]}"', allow_blank=True)
                        ws.add_data_validation(dv_enums[opt])
                    dv_enums[opt].add(ref)
                elif fmt == "d":
                    dv_date.add(ref)
                    ac.number_format = "DD.MM.YYYY"
                elif fmt == "n":
                    dv_num.add(ref)
                    ac.number_format = "#,##0.00" if unit not in ("Jahr", "Anzahl") else "0"
                elif fmt == "p":
                    ac.number_format = "0.0%"
                elif fmt == "w":
                    dv_whole.add(ref)
                    ac.number_format = "0"
                elif fmt == "tm":
                    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
                    ac.alignment = WRAP
                r += 1
        last_q = r - 1
        ws["G2"] = f"=COUNTA(C{first_q}:C{last_q})/COUNTA(A{first_q}:A{last_q})"
        ws["G2"].number_format = "0%"
        ws.conditional_formatting.add(
            f"C{first_q}:C{last_q}",
            FormulaRule(formula=[f'AND($A{first_q}<>"",ISBLANK(C{first_q}))'], fill=PatternFill("solid", fgColor="FCE4D6")),
        )
        summary.cell(row=2 + s_idx, column=1, value=sheet_name)
        summary.cell(row=2 + s_idx, column=2, value=f"='{sheet_name}'!G2").number_format = "0%"
    assert total == 100, total
    return _save(wb, path)


# —— CBAM communication（Belegt 样例版面）——

CBAM_ROUTE_OPTIONS = ["Blast furnace route", "Electric arc furnace", "DRI-EAF", "Other"]
CBAM_METHOD_OPTIONS = ["Actual (verified)", "Default value", "Other"]
CBAM_INST_LABELS = ["Installation identifier", "Name of the installation", "Operator name", "Country", "Address",
                    "UN/LOCODE", "Reporting period start", "Reporting period end", "Contact e-mail"]
CBAM_PRODUCT_COLS = ["No.", "CN code", "Goods description", "Production route", "SEE (direct)", "SEE (indirect)",
                     "Specific electricity consumption", "Determination method"]
CBAM_PROCESS_COLS = ["No.", "Production process", "Production route", "CN code", "Activity level"]


def build_cbam_template(path: str | Path, rows: int = 10) -> Path:
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    ins["A1"] = "CBAM communication template for installations – operator to importer (Belegt sample layout)"
    ins["A1"].font = TITLE_FONT
    ins["A3"] = ("Structure follows the EU Commission communication template (sheets A_InstData, D_Processes, "
                 "Summary_Products). Cell positions of the official template must be verified for each release.")
    a = wb.create_sheet("A_InstData")
    a["A1"] = "A. Installation data"
    a["A1"].font = TITLE_FONT
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
    a.add_data_validation(dv_date)
    for i, label in enumerate(CBAM_INST_LABELS, start=3):
        a.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        a.cell(row=i, column=1, value=label).font = Font(bold=True)
        c = a.cell(row=i, column=3)
        c.fill, c.border = INPUT_FILL, BORDER
        if "period" in label:
            dv_date.add(c.coordinate)
            c.number_format = "YYYY-MM-DD"
    a.column_dimensions["A"].width = 28
    a.column_dimensions["C"].width = 40

    def table(name: str, title: str, cols: list[str], dv_specs: dict[str, DataValidation], fmts: dict[str, str]) -> None:
        ws = wb.create_sheet(name)
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        ws["A2"] = "Units: SEE in tCO2e/t; specific electricity consumption in MWh/t; activity level in t."
        for c, h in enumerate(cols, start=1):
            cell = ws.cell(row=3, column=c, value=h)
            cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
            ws.column_dimensions[cell.column_letter].width = max(14, len(h) + 4)
        for dv in dv_specs.values():
            ws.add_data_validation(dv)
        for i in range(1, rows + 1):
            r = 3 + i
            ws.cell(row=r, column=1, value=i).border = BORDER
            for c, h in enumerate(cols[1:], start=2):
                cell = ws.cell(row=r, column=c)
                cell.fill, cell.border = INPUT_FILL, BORDER
                if h in dv_specs:
                    dv_specs[h].add(cell.coordinate)
                if h in fmts:
                    cell.number_format = fmts[h]

    num = lambda: DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)  # noqa: E731
    table("Summary_Products", "Summary of specific embedded emissions per product", CBAM_PRODUCT_COLS, {
        "Production route": DataValidation(type="list", formula1='"' + ",".join(CBAM_ROUTE_OPTIONS) + '"', allow_blank=True),
        "Determination method": DataValidation(type="list", formula1='"' + ",".join(CBAM_METHOD_OPTIONS) + '"', allow_blank=True),
        "SEE (direct)": num(), "SEE (indirect)": num(), "Specific electricity consumption": num(),
    }, {"SEE (direct)": "0.000000", "SEE (indirect)": "0.000000", "Specific electricity consumption": "0.000",
        "CN code": "@"})
    table("D_Processes", "D. Production processes", CBAM_PROCESS_COLS, {
        "Production route": DataValidation(type="list", formula1='"' + ",".join(CBAM_ROUTE_OPTIONS) + '"', allow_blank=True),
        "Activity level": num(),
    }, {"Activity level": "#,##0.000", "CN code": "@"})
    return _save(wb, path)


# —— CMRT（简化版面）——

CMRT_METALS = ["Tantalum", "Tin", "Gold", "Tungsten"]
CMRT_DECL_QUESTIONS = [
    ("Q1", "3TG intentionally added or used in the product(s)?"),
    ("Q2", "3TG necessary to the production and contained in the finished product(s)?"),
    ("Q3", "Any 3TG originating from covered countries?"),
    ("Q4", "Any 3TG coming from recycled or scrap sources?"),
    ("Q5", "Have all smelters in your supply chain been identified?"),
    ("Q6", "Have all applicable smelters been reported in this declaration?"),
    ("Q7", "Have you identified all smelters for the declared scope?"),
]
CMRT_COMPANY_QUESTIONS = [
    ("A", "Do you have a responsible minerals sourcing policy?"),
    ("B", "Is your policy publicly available on your website? (provide URL)"),
    ("C", "Do you require your direct suppliers to source from validated smelters?"),
    ("D", "Have you implemented due diligence measures for responsible sourcing?"),
    ("E", "Do you review due diligence information received from suppliers against your expectations?"),
    ("F", "Does your review process include corrective action management?"),
    ("G", "Are you subject to an annual conflict minerals disclosure requirement?"),
    ("H", "Do you collect smelter information from your suppliers?"),
]
CMRT_INFO_LABELS = ["Company Name", "Declaration Scope", "Authorized Company Representative", "Contact e-mail",
                    "Date of Completion"]
CMRT_SMELTER_COLS = ["Metal", "Smelter Name", "Smelter Country", "Smelter Identification", "Comments"]
CMRT_PRODUCT_COLS = ["Manufacturer's Product Number", "Manufacturer's Product Name", "Comments"]


def build_cmrt_template(path: str | Path, smelter_rows: int = 20, product_rows: int = 10) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Declaration"
    ws["A1"] = "Conflict Minerals Reporting Template (CMRT) – simplified export, structure based on RMI CMRT 6.5"
    ws["A1"].font = TITLE_FONT
    for i, label in enumerate(CMRT_INFO_LABELS, start=3):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=i, column=3)
        c.fill, c.border = INPUT_FILL, BORDER
    scope = DataValidation(type="list", formula1='"Company,Product (or List of Products),User defined"', allow_blank=True)
    ws.add_data_validation(scope)
    scope.add("C4")
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
    ws.add_data_validation(dv_date)
    dv_date.add("C7")
    ws["C7"].number_format = "DD-MMM-YYYY"
    hr = 9
    for c, h in enumerate(["Nr", "Question", *CMRT_METALS], start=1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
    ynu = DataValidation(type="list", formula1='"Yes,No,Unknown"', allow_blank=True)
    ws.add_data_validation(ynu)
    for i, (nid, text) in enumerate(CMRT_DECL_QUESTIONS, start=hr + 1):
        ws.cell(row=i, column=1, value=nid)
        ws.cell(row=i, column=2, value=text).alignment = WRAP
        for c in range(3, 3 + len(CMRT_METALS)):
            cell = ws.cell(row=i, column=c)
            cell.fill, cell.border = INPUT_FILL, BORDER
            ynu.add(cell.coordinate)
    hr2 = hr + len(CMRT_DECL_QUESTIONS) + 2
    for c, h in enumerate(["Nr", "Company question", "Answer", "Comments"], start=1):
        cell = ws.cell(row=hr2, column=c, value=h)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
    yn = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(yn)
    for i, (nid, text) in enumerate(CMRT_COMPANY_QUESTIONS, start=hr2 + 1):
        ws.cell(row=i, column=1, value=nid)
        ws.cell(row=i, column=2, value=text).alignment = WRAP
        cell = ws.cell(row=i, column=3)
        cell.fill, cell.border = INPUT_FILL, BORDER
        if nid != "B":
            yn.add(cell.coordinate)
    ws.column_dimensions["B"].width = 70
    for col in "CDEF":
        ws.column_dimensions[col].width = 14

    def table(name: str, cols: list[str], n: int, dvs: dict[str, DataValidation]) -> None:
        t = wb.create_sheet(name)
        t["A1"] = f"{name} (CMRT simplified)"
        t["A1"].font = TITLE_FONT
        for c, h in enumerate(cols, start=1):
            cell = t.cell(row=3, column=c, value=h)
            cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
            t.column_dimensions[cell.column_letter].width = max(16, len(h) + 4)
        for dv in dvs.values():
            t.add_data_validation(dv)
        for r in range(4, 4 + n):
            for c, h in enumerate(cols, start=1):
                cell = t.cell(row=r, column=c)
                cell.fill, cell.border = INPUT_FILL, BORDER
                if h in dvs:
                    dvs[h].add(cell.coordinate)

    table("Smelter List", CMRT_SMELTER_COLS, smelter_rows,
          {"Metal": DataValidation(type="list", formula1='"' + ",".join(CMRT_METALS) + '"', allow_blank=True)})
    table("Product List", CMRT_PRODUCT_COLS, product_rows, {})
    return _save(wb, path)


# —— CSV ——

CSV_QUESTIONS = [
    ("1", "Firmenname", "", ""),
    ("2", "Anzahl Mitarbeitende zum 31.12.2025", "", "Anzahl"),
    ("3", "Scope-1-Emissionen im Berichtsjahr 2025", "", "tCO2e"),
    ("4", "Scope-2-Emissionen (marktbasiert) im Berichtsjahr 2025", "", "tCO2e"),
    ("5", "Stromverbrauch im Berichtsjahr 2025", "", "MWh"),
    ("6", "Anteil erneuerbarer Energien am Stromverbrauch", "", "%"),
    ("7", "Ist Ihr Standort nach ISO 14001 zertifiziert?", "", ""),
    ("8", "ISO 14001: gültig bis", "", ""),
    ("9", "Enthält das Produkt SVHC-Stoffe über 0,1 Massen-%?", "", ""),
    ("10", "PCF des gelieferten Produkts (cradle-to-gate)", "", "kgCO2e/kg"),
    ("11", "Wasserentnahme gesamt im Berichtsjahr 2025", "", "m³"),
    ("12", "Abfallaufkommen gesamt im Berichtsjahr 2025", "", "t"),
    ("13", "Liefern Sie Waren, die unter den CBAM fallen?", "", ""),
    ("14", "Spezifische direkte graue Emissionen (SEE direkt)", "", "tCO2e/t"),
    ("15", "Link zum aktuellen Nachhaltigkeitsbericht", "", ""),
]


def build_csv_questionnaire(path: str | Path) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";", lineterminator="\r\n")
        w.writerow(["Nr.", "Frage", "Antwort", "Einheit", "Kommentar"])
        for nid, text, ans, unit in CSV_QUESTIONS:
            w.writerow([nid, text, ans, unit, ""])
    return p


# —— 示例 PDF（无第三方依赖的最小 PDF 写出）——

def build_sample_pdf(path: str | Path, lines: list[str] | None = None) -> Path:
    lines = lines or [
        "Stadtwerke Musterstadt GmbH",
        "Jahresabrechnung Strom 2025",
        "Abrechnungszeitraum 01.01.2025 - 31.12.2025",
        "Wirkarbeit gesamt: 4.820.000 kWh",
        "davon Oekostrom: 1.234,5 MWh",
        "Erdgas: 2.150 MWh",
        "CO2-Emissionen (standortbezogen): 1.835,2 t CO2e",
        "Rechnungsbetrag: 1.234.567,89 EUR",
    ]

    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    y = 780
    ops = ["BT", "/F1 11 Tf"]
    for i, ln in enumerate(lines):
        ops.append(f"1 0 0 1 72 {y - i * 18} Tm ({esc(ln)}) Tj")
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(bytes(out))
    return p
