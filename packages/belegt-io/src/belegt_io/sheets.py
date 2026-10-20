# Spec: 01 AC-501-1（合并单元格、下拉校验）；03 §3.5 Locator
"""xlsx / csv 的统一只读视图：合并单元格、数据验证（下拉）解析、单元格样式。"""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet


def coord(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def split_coord(c: str) -> tuple[int, int]:
    letters, row = coordinate_from_string(c)
    return row, column_index_from_string(letters)


class DVInfo:
    __slots__ = ("type", "options", "formula", "allow_blank", "operator", "formula2")

    def __init__(self, type_: str | None, options: list[str] | None, formula: str | None,
                 allow_blank: bool = True, operator: str | None = None, formula2: str | None = None):
        self.type = type_
        self.options = options
        self.formula = formula
        self.allow_blank = allow_blank
        self.operator = operator
        self.formula2 = formula2


class SheetView:
    """抽象：行列 1 起。"""

    name: str | None
    max_row: int
    max_col: int

    def value(self, r: int, c: int) -> Any:  # pragma: no cover - interface
        raise NotImplementedError

    def merged_bounds(self, r: int, c: int) -> tuple[int, int, int, int] | None:
        return None

    def dv(self, r: int, c: int) -> DVInfo | None:
        return None

    def bold(self, r: int, c: int) -> bool:
        return False

    def number_format(self, r: int, c: int) -> str | None:
        return None

    def is_formula(self, r: int, c: int) -> bool:
        v = self.value(r, c)
        return isinstance(v, str) and v.startswith("=")

    def text(self, r: int, c: int) -> str:
        v = self.value(r, c)
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v).strip()

    def top_left(self, r: int, c: int) -> tuple[int, int]:
        b = self.merged_bounds(r, c)
        return (b[0], b[1]) if b else (r, c)


class XlsxSheetView(SheetView):
    def __init__(self, ws: Worksheet, wb: Any):
        self.ws = ws
        self.wb = wb
        self.name = ws.title
        self.max_row = ws.max_row or 0
        self.max_col = ws.max_column or 0
        self._merged: dict[tuple[int, int], tuple[int, int, int, int]] = {}
        for mr in ws.merged_cells.ranges:
            min_c, min_r, max_c, max_r = mr.bounds
            for rr in range(min_r, max_r + 1):
                for cc in range(min_c, max_c + 1):
                    self._merged[(rr, cc)] = (min_r, min_c, max_r, max_c)
        self._dvs = list(ws.data_validations.dataValidation) if ws.data_validations else []
        self._dv_cache: dict[tuple[int, int], DVInfo | None] = {}

    def value(self, r: int, c: int) -> Any:
        return self.ws.cell(row=r, column=c).value

    def merged_bounds(self, r: int, c: int) -> tuple[int, int, int, int] | None:
        return self._merged.get((r, c))

    def bold(self, r: int, c: int) -> bool:
        f = self.ws.cell(row=r, column=c).font
        return bool(f and f.b)

    def number_format(self, r: int, c: int) -> str | None:
        return self.ws.cell(row=r, column=c).number_format

    def dv(self, r: int, c: int) -> DVInfo | None:
        key = (r, c)
        if key in self._dv_cache:
            return self._dv_cache[key]
        ref = coord(r, c)
        info = None
        for dv in self._dvs:
            if ref in dv.sqref:
                opts = None
                if dv.type == "list":
                    opts = resolve_list_formula(dv.formula1, self.wb, self.ws)
                info = DVInfo(dv.type, opts, dv.formula1, bool(dv.allow_blank), dv.operator, dv.formula2)
                break
        self._dv_cache[key] = info
        return info


def _clean_opt(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def resolve_list_formula(formula: str | None, wb: Any, ws: Worksheet | None, _depth: int = 0) -> list[str] | None:
    """解析数据验证 list 的 formula1：字面量 `"a,b"`、区域 `$A$1:$A$3` / `Sheet!$A$1:$A$3`、定义名称。"""
    if not formula or _depth > 3:
        return None
    f = formula.strip()
    if f.startswith("="):
        f = f[1:]
    if f.startswith('"') and f.endswith('"'):
        return [o for o in (x.strip() for x in f[1:-1].split(",")) if o]
    # 定义名称
    if re.fullmatch(r"[A-Za-z_][\w.]*", f) and not re.fullmatch(r"[A-Z]{1,3}\d+", f):
        dn = None
        try:
            dn = wb.defined_names.get(f)
        except Exception:  # pragma: no cover
            dn = None
        if dn is None and ws is not None:
            try:
                dn = ws.defined_names.get(f)
            except Exception:  # pragma: no cover
                dn = None
        if dn is not None:
            return resolve_list_formula(dn.attr_text, wb, ws, _depth + 1)
        return None
    sheet_name = None
    rng = f
    if "!" in f:
        sheet_name, _, rng = f.rpartition("!")
        sheet_name = sheet_name.strip("'").replace("''", "'")
    target = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else ws
    if target is None:
        return None
    try:
        min_c, min_r, max_c, max_r = range_boundaries(rng.replace("$", ""))
    except (ValueError, TypeError):
        return None
    out: list[str] = []
    for r in range(min_r, (max_r or min_r) + 1):
        for c in range(min_c, (max_c or min_c) + 1):
            s = _clean_opt(target.cell(row=r, column=c).value)
            if s:
                out.append(s)
    return out


class GridSheetView(SheetView):
    """CSV：二维字符串表。"""

    def __init__(self, rows: list[list[str]], name: str | None = None):
        self.rows = rows
        self.name = name
        self.max_row = len(rows)
        self.max_col = max((len(r) for r in rows), default=0)

    def value(self, r: int, c: int) -> Any:
        if 1 <= r <= len(self.rows) and 1 <= c <= len(self.rows[r - 1]):
            v = self.rows[r - 1][c - 1]
            return v if v != "" else None
        return None


def read_csv_rows(path: str | Path) -> tuple[list[list[str]], str, str]:
    """读取 CSV：编码 utf-8-sig → cp1252 回退；分隔符嗅探（; , \\t |）。返回 (rows, encoding, delimiter)。"""
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        raise ValueError("cannot decode CSV")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if sample.count(";") > sample.count(",") else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    return rows, enc, delim
