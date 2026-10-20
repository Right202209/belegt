# Spec: 01 FR-501（AC-501-1/2/3）；03 §3.3 Question.locator；04 §2.11 filler-file（确定性回填）；02 §3.2（导出舍入）
"""原格式回填 `fill_back`：按 locator 把答案写回原 xlsx/csv。

保证：
- 只写答案单元格（以解析结果中的 `QuestionDraft.locator` 为白名单），不碰其它单元格；含公式的单元格拒写；
- 保留格式、公式、合并单元格、数据验证、条件格式（openpyxl 往返）；
- 日期/数字/枚举按目标单元格做格式适配，写入前校验枚举合法性与数字校验规则；
- 生成回填报告（变更单元格列表，AC-501-2）与 `export_input_hash`（AC-501-3）；
- 输出字节确定：同一源文件 + 同一答案集 → 同一 sha256（固定 zip 时间戳与 core.xml modified）。
"""
from __future__ import annotations

import csv
import io
import re
import shutil
import tempfile
import zipfile
from datetime import date, datetime
from decimal import ROUND_HALF_EVEN, Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Mapping

from .hashing import canonical_hash, file_sha256
from .models import (
    BelegtIOError, CellChange, FillReport, IOErrorCode, Locator, ParsedQuestionnaire, QuestionDraft,
)
from .numbers import parse_number, to_dec_str, to_decimal
from .sheets import read_csv_rows, split_coord

YES_WORDS = {"ja", "yes", "true", "wahr", "j", "y", "oui", "sí", "si", "tak"}
NO_WORDS = {"nein", "no", "false", "falsch", "n", "non", "nie"}
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)


class _Reject(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def _unwrap(v: Any) -> tuple[Any, int | None]:
    if isinstance(v, dict) and "value" in v:
        return v["value"], v.get("round")
    return v, None


def _as_date(v: Any) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise _Reject(IOErrorCode.TYPE_MISMATCH, f"not a date: {v!r}")


def _as_decimal(v: Any, lang: str | None) -> Decimal:
    if isinstance(v, bool):
        raise _Reject(IOErrorCode.TYPE_MISMATCH, "bool given for number cell")
    try:
        return to_decimal(v)
    except (InvalidOperation, TypeError, ValueError):
        pass
    try:
        pn = parse_number(str(v), lang=lang)
    except ValueError as e:
        raise _Reject(IOErrorCode.TYPE_MISMATCH, f"not a number: {v!r}") from e
    if pn.needs_confirmation:
        raise _Reject(IOErrorCode.TYPE_MISMATCH, f"ambiguous number {v!r} ({','.join(pn.reasons)})")
    return pn.value


def _match_option(v: Any, options: list[str]) -> str:
    if isinstance(v, bool):
        words = YES_WORDS if v else NO_WORDS
        for o in options:
            if o.strip().lower() in words:
                return o
        raise _Reject(IOErrorCode.ENUM_INVALID, f"no {'yes' if v else 'no'} option in {options}")
    s = str(v).strip()
    for o in options:
        if o == s:
            return o
    for o in options:
        if o.strip().lower() == s.lower():
            return o
    # 数字型选项（如 "1"/"2"）
    try:
        d = to_decimal(v)
        for o in options:
            try:
                if Decimal(o) == d:
                    return o
            except InvalidOperation:
                continue
    except (InvalidOperation, TypeError, ValueError):
        pass
    raise _Reject(IOErrorCode.ENUM_INVALID, f"{s!r} not in allowed options {options}")


def _check_dv_number(d: Decimal, dv: dict[str, Any] | None, cell_dv: Any) -> None:
    if cell_dv is None:
        return
    if cell_dv.type == "whole" and d != d.to_integral_value():
        raise _Reject(IOErrorCode.TYPE_MISMATCH, f"{d} is not a whole number (data validation)")
    op = cell_dv.operator or "between"
    try:
        f1 = Decimal(str(cell_dv.formula)) if cell_dv.formula not in (None, "") else None
        f2 = Decimal(str(cell_dv.formula2)) if cell_dv.formula2 not in (None, "") else None
    except InvalidOperation:
        return  # 公式型边界不做本地求值
    ok = True
    if op == "between" and f1 is not None and f2 is not None:
        ok = f1 <= d <= f2
    elif op == "notBetween" and f1 is not None and f2 is not None:
        ok = not (f1 <= d <= f2)
    elif op == "greaterThanOrEqual" and f1 is not None:
        ok = d >= f1
    elif op == "greaterThan" and f1 is not None:
        ok = d > f1
    elif op == "lessThanOrEqual" and f1 is not None:
        ok = d <= f1
    elif op == "lessThan" and f1 is not None:
        ok = d < f1
    elif op == "equal" and f1 is not None:
        ok = d == f1
    if not ok:
        raise _Reject(IOErrorCode.TYPE_MISMATCH, f"{d} violates data validation {op} {f1} {f2 or ''}".strip())


def _round(d: Decimal, digits: int | None) -> Decimal:
    if digits is None:
        return d
    return d.quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_EVEN)


def adapt_value(q: QuestionDraft | None, value: Any, *, number_format: str | None = None, cell_dv: Any = None,
                lang: str | None = None, round_digits: int | None = None) -> tuple[Any, str | None, list[str]]:
    """把答案值适配为写入单元格的 Python 值。返回 (值, 需设置的 number_format 或 None, 说明)。"""
    notes: list[str] = []
    fmt = q.answer_format if q else "text"
    options = (q.options if q else None) or (cell_dv.options if cell_dv is not None and cell_dv.type == "list" else None)
    lang = lang or (q.language if q else None)
    if options:
        return _match_option(value, options), None, notes
    if fmt == "bool" or isinstance(value, bool):
        if isinstance(value, bool):
            yes, no = ("Ja", "Nein") if (lang or "").startswith("de") else ("Yes", "No")
            return (yes if value else no), None, notes
    if fmt == "date" or (cell_dv is not None and cell_dv.type == "date"):
        d = _as_date(value)
        new_fmt = None
        if not number_format or number_format == "General":
            new_fmt = "DD.MM.YYYY" if (lang or "").startswith("de") else "YYYY-MM-DD"
        return datetime(d.year, d.month, d.day), new_fmt, notes
    if fmt == "number" or (cell_dv is not None and cell_dv.type in ("decimal", "whole")):
        d = _as_decimal(value, lang)
        d = _round(d, round_digits)
        if number_format and "%" in number_format:
            d = d / Decimal(100)  # 字典存 0–100（02 §3.1），百分比格式单元格存比值
            notes.append("percent_cell: value/100")
        _check_dv_number(d, None, cell_dv)  # 数据验证作用于单元格实际存储值
        if d == d.to_integral_value() and abs(d) < Decimal(10) ** 15:
            return int(d), None, notes
        return d, None, notes
    if isinstance(value, (date, datetime)):
        return value.isoformat(), None, notes
    if isinstance(value, (Decimal, int)) and not isinstance(value, bool):
        d = _round(Decimal(value), round_digits)  # 题型误判为 text 时仍按数字写，保留 Excel 数值类型
        return (int(d) if d == d.to_integral_value() else d), None, notes
    return value if isinstance(value, str) else str(value), None, notes


def _values_equal(a: Any, b: Any) -> bool:
    if a is None or b is None:
        return a is b
    try:
        if isinstance(a, (int, float, Decimal)) and isinstance(b, (int, float, Decimal)) \
                and not isinstance(a, bool) and not isinstance(b, bool):
            return Decimal(str(a)) == Decimal(str(b))
    except InvalidOperation:  # pragma: no cover
        return False
    if isinstance(a, datetime) and isinstance(b, datetime):
        return a == b
    return a == b


def _answers_canon(answers: Mapping[Any, Any]) -> list[list[Any]]:
    out = []
    for k, v in answers.items():
        loc = k if isinstance(k, Locator) else Locator.parse(str(k))
        out.append([loc.sheet_cell or str(loc), _unwrap(v)[0], _unwrap(v)[1]])
    out.sort(key=lambda x: str(x[0]))
    return out


def export_input_hash(src_sha256: str, answers: Mapping[Any, Any]) -> str:
    """AC-501-3：源文件哈希 + 规范化答案集 → `sha256:`。"""
    return canonical_hash({"src": src_sha256, "answers": _answers_canon(answers)})


def _deterministic_rezip(path: Path, modified_iso: str) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                txt = data.decode("utf-8")
                txt = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rf"\g<1>{modified_iso}\g<2>", txt)
                data = txt.encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            zout.writestr(zi, data)
    tmp.replace(path)


def _resolve_key(key: Any, src_sha: str, default_sheet: str | None) -> Locator:
    loc = key if isinstance(key, Locator) else Locator.parse(str(key))
    if loc.cell is None:
        raise _Reject(IOErrorCode.LOCATOR_INVALID, f"locator has no cell: {key!r}")
    if loc.blob_sha256 and loc.blob_sha256 != src_sha:
        raise _Reject(IOErrorCode.HASH_MISMATCH, f"locator blob {loc.blob_sha256[:19]}… != source {src_sha[:19]}…")
    if loc.sheet is None and default_sheet is not None:
        loc = loc.model_copy(update={"sheet": default_sheet})
    return loc


def fill_back(src_path: str | Path, answers: Mapping[Any, Any], out_path: str | Path, *,
              parsed: ParsedQuestionnaire | None = None, strict: bool = True, lang: str | None = None,
              modified_at: str | None = None) -> FillReport:
    """把 `answers`（locator → 值）写回原格式文件。

    - `answers` 键：`Locator`、`"sha256:<hex>#Sheet!B12"`、`"Sheet!B12"`、CSV 为 `"C5"`；
      值：Decimal/int/str/bool/date/datetime，或 `{"value": v, "round": n}`（02 §3.2 导出舍入）。
    - `strict=True`：只允许写解析出的答案单元格；其它单元格记 `IO_NOT_ANSWER_CELL` 拒绝。
    - 被拒的条目不写入，其余照常写入；调用方据 `report.ok` 决定是否放行。
    """
    src = Path(src_path)
    out = Path(out_path)
    src_sha = file_sha256(src)
    if parsed is None:
        from .parse import parse_questionnaire

        parsed = parse_questionnaire(src)
    if parsed.blob_sha256 != src_sha:
        raise BelegtIOError(IOErrorCode.HASH_MISMATCH, "parsed questionnaire does not belong to src_path", str(src))
    index: dict[tuple[str | None, str], QuestionDraft] = {
        (q.locator.sheet, q.locator.cell or ""): q for q in parsed.questions
    }
    report = FillReport(src_path=str(src), out_path=str(out), src_sha256=src_sha,
                        export_input_hash=export_input_hash(src_sha, answers))
    lang = lang or (parsed.language if parsed.language != "und" else None)
    if src.suffix.lower() == ".csv":
        _fill_csv(src, out, answers, index, report, strict, lang, src_sha)
    else:
        _fill_xlsx(src, out, answers, index, report, strict, lang, src_sha, modified_at)
    report.out_sha256 = file_sha256(out)
    return report


def _fill_xlsx(src: Path, out: Path, answers: Mapping[Any, Any], index: dict, report: FillReport, strict: bool,
               lang: str | None, src_sha: str, modified_at: str | None) -> None:
    from openpyxl import load_workbook

    from .sheets import XlsxSheetView

    with zipfile.ZipFile(src) as z:
        names = z.namelist()
    if any(n.startswith(("xl/media/", "xl/charts/")) for n in names):
        report.warnings.append("source contains images/charts; openpyxl may not round-trip them — verify output")
    if any("extLst" in n for n in names):  # pragma: no cover - 罕见
        report.warnings.append("source contains extension lists")
    keep_vba = src.suffix.lower() == ".xlsm"
    wb = load_workbook(src, data_only=False, keep_vba=keep_vba)
    orig_modified = wb.properties.modified
    default_sheet = wb.sheetnames[0] if len(wb.sheetnames) == 1 else None
    views: dict[str, XlsxSheetView] = {}
    for key, raw in answers.items():
        value, round_digits = _unwrap(raw)
        try:
            loc = _resolve_key(key, src_sha, default_sheet)
        except _Reject as e:
            report.changes.append(CellChange(locator=str(key), sheet=None, cell=str(key), status="rejected",
                                             code=e.code, message=e.message, new_value=value))
            continue
        if loc.sheet not in wb.sheetnames:
            report.changes.append(CellChange(locator=str(key), sheet=loc.sheet, cell=loc.cell or "", status="rejected",
                                             code=IOErrorCode.LOCATOR_INVALID, message="sheet not found",
                                             new_value=value))
            continue
        ws = wb[loc.sheet]
        view = views.setdefault(loc.sheet, XlsxSheetView(ws, wb))
        r, c = split_coord(loc.cell)  # type: ignore[arg-type]
        tr, tc = view.top_left(r, c)
        cell = ws.cell(row=tr, column=tc)
        q = index.get((loc.sheet, cell.coordinate))
        ch = CellChange(locator=f"{src_sha}#{loc.sheet}!{cell.coordinate}", sheet=loc.sheet, cell=cell.coordinate,
                        old_value=cell.value, new_value=value, number_format=cell.number_format,
                        question_ordinal=q.ordinal if q else None, native_id=q.native_id if q else None)
        report.changes.append(ch)
        if value is None:
            ch.status, ch.message = "skipped", "value is None (empty answer semantics live in Answer.status)"
            continue
        if (tr, tc) != (r, c):
            ch.message = f"merged cell → top-left {cell.coordinate}"
        if strict and q is None:
            ch.status, ch.code, ch.message = "rejected", IOErrorCode.NOT_ANSWER_CELL, "not an answer cell of this questionnaire"
            continue
        if isinstance(cell.value, str) and cell.value.startswith("="):
            ch.status, ch.code, ch.message = "rejected", IOErrorCode.FORMULA_CELL, "target cell contains a formula"
            continue
        try:
            new, new_fmt, notes = adapt_value(q, value, number_format=cell.number_format, cell_dv=view.dv(tr, tc),
                                              lang=lang, round_digits=round_digits)
        except _Reject as e:
            ch.status, ch.code, ch.message = "rejected", e.code, e.message
            continue
        if _values_equal(cell.value, new):
            ch.status, ch.new_value = "unchanged", new
            continue
        cell.value = new
        if new_fmt:
            cell.number_format = new_fmt
        ch.new_value = new
        ch.number_format = cell.number_format
        if notes:
            ch.message = "; ".join(filter(None, [ch.message, *notes]))
    tmpdir = Path(tempfile.mkdtemp(prefix="belegt-io-"))
    try:
        tmp = tmpdir / out.name
        wb.save(tmp)
        mod = modified_at or (orig_modified.strftime("%Y-%m-%dT%H:%M:%SZ") if orig_modified else "1980-01-01T00:00:00Z")
        _deterministic_rezip(tmp, mod)
        out.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(tmp), out)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def _fmt_csv_value(v: Any, lang: str | None) -> str:
    de = (lang or "").startswith("de")
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y") if de else v.date().isoformat()
    if isinstance(v, (int, Decimal)) and not isinstance(v, bool):
        s = to_dec_str(Decimal(v))
        return s.replace(".", ",") if de else s
    return str(v)


def _fill_csv(src: Path, out: Path, answers: Mapping[Any, Any], index: dict, report: FillReport, strict: bool,
              lang: str | None, src_sha: str) -> None:
    rows, enc, delim = read_csv_rows(src)
    for key, raw in answers.items():
        value, round_digits = _unwrap(raw)
        try:
            loc = _resolve_key(key, src_sha, None)
        except _Reject as e:
            report.changes.append(CellChange(locator=str(key), sheet=None, cell=str(key), status="rejected",
                                             code=e.code, message=e.message, new_value=value))
            continue
        r, c = split_coord(loc.cell)  # type: ignore[arg-type]
        while len(rows) < r:
            rows.append([])
        row = rows[r - 1]
        while len(row) < c:
            row.append("")
        q = index.get((None, loc.cell))
        ch = CellChange(locator=f"{src_sha}#{loc.cell}", sheet=None, cell=loc.cell or "", old_value=row[c - 1] or None,
                        new_value=value, question_ordinal=q.ordinal if q else None,
                        native_id=q.native_id if q else None)
        report.changes.append(ch)
        if value is None:
            ch.status = "skipped"
            continue
        if strict and q is None:
            ch.status, ch.code, ch.message = "rejected", IOErrorCode.NOT_ANSWER_CELL, "not an answer cell of this questionnaire"
            continue
        try:
            new, _, notes = adapt_value(q, value, lang=lang, round_digits=round_digits)
        except _Reject as e:
            ch.status, ch.code, ch.message = "rejected", e.code, e.message
            continue
        s = _fmt_csv_value(new, lang)
        if row[c - 1] == s:
            ch.status, ch.new_value = "unchanged", s
            continue
        row[c - 1] = s
        ch.new_value = s
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delim, lineterminator="\r\n")
    w.writerows(rows)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(buf.getvalue().encode("utf-8-sig" if enc == "utf-8-sig" else enc))
